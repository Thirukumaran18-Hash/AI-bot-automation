import csv
import io
import re
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..deps import require_roles
from ..models import Assignment, Evaluation, Intern, Meeting, Task, User, WorkflowEvent
from ..schemas import InternCreate, InternUpdate
from ..security import create_temp_password, hash_password
from ..serializers import intern_dict
from ..services import autopilot
from ..services.email_service import send_email
from ..services.workflow import record_event

router = APIRouter(prefix="/interns", tags=["interns"])

COLUMN_ALIASES = {
    "name": {"name", "fullname", "candidate", "candidatename", "intern", "studentname"},
    "email": {"email", "mail", "mailid", "emailid", "emailaddress"},
    "role": {
        "role",
        "domain",
        "position",
        "designation",
        "jobrole",
        "track",
        "technology",
        "assignedrole",
        "jobtitle",
        "rolename",
    },
    "phone": {
        "phone",
        "phonenumber",
        "mobile",
        "mobilenumber",
        "contact",
        "contactnumber",
        "phoneno",
    },
    "batch": {"batch", "cohort", "group", "company", "companyname"},
    "skills": {"skills", "skillset", "department", "degreebranch", "branch", "degree"},
    "mentor_email": {"mentoremail", "mentor"},
    "password": {"password", "temppassword"},
    "initial_task_title": {"initialtasktitle", "tasktitle", "firsttasktitle", "firsttask"},
    "initial_task_id": {"initialtaskid", "taskid"},
}


def _normalise_key(value: str) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\([^)]*\)", "", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def _canonical_row(row: dict) -> dict:
    result: dict[str, str] = {}
    for key, value in row.items():
        if key is None:
            continue
        normalised = _normalise_key(key)
        for canonical, aliases in COLUMN_ALIASES.items():
            if normalised in {_normalise_key(alias) for alias in aliases}:
                result[canonical] = "" if value is None else str(value).strip()
                break
    return result


def _read_upload(raw: bytes, filename: str | None) -> list[dict]:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(
                status_code=500,
                detail="Excel support is not installed. Run: pip install openpyxl",
            ) from exc
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Could not read the Excel file: {exc}") from exc
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = list(rows[0])
        payload: list[dict] = []
        for values in rows[1:]:
            if not any(value not in (None, "") for value in values):
                continue
            payload.append(
                {
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                    if headers[index] is not None
                }
            )
        return payload

    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded") from exc
    return [
        {key: value for key, value in row.items() if key}
        for row in csv.DictReader(io.StringIO(content))
    ]


def _mentor_or_404(db: Session, mentor_id: int | None) -> User | None:
    if mentor_id is None:
        return None
    mentor = db.get(User, mentor_id)
    if not mentor or mentor.role != "mentor" or not mentor.active:
        raise HTTPException(status_code=400, detail="Selected mentor is not available")
    return mentor


def _visible_interns(db: Session, user: User):
    query = db.query(Intern)
    if user.role == "mentor":
        query = query.filter(Intern.mentor_id == user.id)
    return query


@router.get("")
def list_interns(
    batch: str | None = None,
    domain: str | None = None,
    search: str | None = None,
    active: bool | None = None,
    user: User = Depends(require_roles("admin", "mentor")),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = _visible_interns(db, user).join(User, Intern.user_id == User.id)
    if batch:
        query = query.filter(Intern.batch == batch)
    if domain:
        query = query.filter(Intern.domain == domain)
    if active is not None:
        query = query.filter(Intern.active.is_(active))
    if search:
        pattern = f"%{search.lower()}%"
        query = query.filter(
            or_(
                User.name.ilike(pattern),
                User.email.ilike(pattern),
                Intern.skills.ilike(pattern),
            )
        )
    return [intern_dict(intern) for intern in query.order_by(User.name).all()]


@router.get("/facets")
def intern_facets(
    user: User = Depends(require_roles("admin", "mentor")),
    db: Session = Depends(get_db),
) -> dict:
    query = _visible_interns(db, user)
    batches = sorted({row[0] for row in query.with_entities(Intern.batch).distinct().all() if row[0]})
    domains = sorted({row[0] for row in query.with_entities(Intern.domain).distinct().all() if row[0]})
    mentors = (
        db.query(User)
        .filter(User.role == "mentor", User.active.is_(True))
        .order_by(User.name)
        .all()
    )
    return {
        "batches": batches,
        "domains": domains,
        "mentors": [{"id": m.id, "name": m.name, "email": m.email} for m in mentors],
    }


@router.post("", status_code=status.HTTP_201_CREATED)
def create_intern(
    payload: InternCreate,
    user: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    existing = db.query(User).filter(User.email == payload.email.lower()).first()
    if existing:
        raise HTTPException(status_code=409, detail="A user with this email already exists")
    mentor = _mentor_or_404(db, payload.mentor_id)

    password = payload.password or create_temp_password()
    account = User(
        email=payload.email.lower(),
        name=payload.name,
        role="intern",
        password_hash=hash_password(password),
    )
    db.add(account)
    db.flush()

    intern = Intern(
        user_id=account.id,
        batch=payload.batch,
        domain=payload.domain,
        skills=payload.skills,
        mentor_id=mentor.id if mentor else None,
    )
    db.add(intern)
    db.flush()
    record_event(
        db,
        "intern",
        intern.id,
        "intern.created",
        new_status="active",
        actor_id=user.id,
        meta={"email": account.email, "batch": intern.batch},
    )
    send_email(
        db,
        recipient=account.email,
        template="welcome",
        context={
            "name": account.name,
            "email": account.email,
            "password": password,
            "batch": intern.batch,
            "domain": intern.domain,
            "portal_url": f"{settings.app_public_url.rstrip('/')}/assignments",
        },
        entity_type="intern",
        entity_id=intern.id,
    )
    db.commit()
    db.refresh(intern)
    return {"intern": intern_dict(intern), "temporary_password": password}


@router.post("/import")
async def import_interns(
    file: UploadFile = File(...),
    launch: bool = Form(default=False),
    user: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="The uploaded file is empty")

    rows = _read_upload(raw, file.filename)
    if not rows:
        raise HTTPException(status_code=400, detail="The uploaded file has no data rows")

    available: set[str] = set()
    for row in rows:
        available.update(_canonical_row(row).keys())
    missing = [field for field in ("name", "email", "role") if field not in available]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=(
                "File needs columns: name, email, role (domain also works). "
                "Optional: phone, batch, skills, mentor_email, password."
            ),
        )

    created, skipped, errors = [], [], []
    for row_number, row in enumerate(rows, start=2):
        normalized = _canonical_row(row)
        name = normalized.get("name", "")
        email = normalized.get("email", "").lower()
        role = normalized.get("role", "")
        if not name or not email or not role:
            errors.append({"row": row_number, "error": "name, email and role are required"})
            continue
        if "@" not in email:
            errors.append({"row": row_number, "error": f"invalid email: {email}"})
            continue
        if db.query(User).filter(func.lower(User.email) == email).first():
            skipped.append({"row": row_number, "email": email, "reason": "already exists"})
            continue

        mentor = None
        mentor_email = normalized.get("mentor_email", "").lower()
        if mentor_email:
            mentor = db.query(User).filter(User.email == mentor_email, User.role == "mentor").first()

        password = normalized.get("password") or create_temp_password()
        account = User(
            email=email,
            name=name,
            role="intern",
            password_hash=hash_password(password),
        )
        db.add(account)
        db.flush()
        intern = Intern(
            user_id=account.id,
            batch=normalized.get("batch") or "General",
            domain=role,
            phone=normalized.get("phone", ""),
            skills=normalized.get("skills") or role,
            mentor_id=mentor.id if mentor else None,
        )
        db.add(intern)
        db.flush()
        record_event(
            db,
            "intern",
            intern.id,
            "intern.imported",
            new_status="active",
            actor_id=user.id,
            meta={"email": email, "role": role, "phone": intern.phone, "source": file.filename},
        )
        send_email(
            db,
            recipient=email,
            template="welcome",
            context={
                "name": name,
                "email": email,
                "password": password,
                "batch": intern.batch,
                "domain": intern.domain,
                "portal_url": f"{settings.app_public_url.rstrip('/')}/assignments",
            },
            entity_type="intern",
            entity_id=intern.id,
        )

        entry = {
            "intern_id": intern.id,
            "name": name,
            "email": email,
            "role": role,
            "phone": intern.phone,
            "temporary_password": password,
        }
        if launch:
            try:
                initial_title = normalized.get("initial_task_title", "")
                if initial_title:
                    autopilot.ensure_custom_first_task(db, role, initial_title, user)
                entry["automation"] = autopilot.launch_intern(
                    db,
                    intern,
                    user,
                    password=password,
                    initial_task_title=initial_title or None,
                )
            except Exception as exc:  # keep the imported candidate even if automation fails
                entry["automation"] = {"status": "error", "reason": str(exc)}
                errors.append({"row": row_number, "error": f"automation failed: {exc}"})
        created.append(entry)

    db.commit()
    launches = [item["automation"] for item in created if item.get("automation")]
    return {
        "created_count": len(created),
        "skipped_count": len(skipped),
        "error_count": len(errors),
        "created": created,
        "skipped": skipped,
        "errors": errors,
        "automation": {
            "requested": launch,
            "launched": sum(1 for item in launches if item.get("status") == "launched"),
            "skipped": sum(1 for item in launches if item.get("status") == "skipped"),
            "failed": sum(1 for item in launches if item.get("status") == "error"),
            "results": launches,
        },
    }


@router.delete("/{intern_id}")
def delete_intern(
    intern_id: int,
    user: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    intern = db.get(Intern, intern_id)
    if not intern:
        raise HTTPException(status_code=404, detail="Intern not found")

    account = intern.user
    previous_status = "active" if intern.active else "inactive"
    name = account.name if account else ""
    email = account.email if account else ""

    for assignment in list(intern.assignments):
        for submission in list(assignment.submissions):
            if submission.file_path:
                try:
                    Path(submission.file_path).unlink(missing_ok=True)
                except OSError:
                    pass

    if account:
        db.query(WorkflowEvent).filter(WorkflowEvent.actor_id == account.id).update(
            {WorkflowEvent.actor_id: None}, synchronize_session=False
        )
        db.query(Evaluation).filter(Evaluation.reviewer_id == account.id).update(
            {Evaluation.reviewer_id: None}, synchronize_session=False
        )
        db.query(Assignment).filter(Assignment.assigned_by == account.id).update(
            {Assignment.assigned_by: None}, synchronize_session=False
        )
        db.query(Meeting).filter(Meeting.created_by == account.id).update(
            {Meeting.created_by: None}, synchronize_session=False
        )
        db.query(Task).filter(Task.created_by == account.id).update(
            {Task.created_by: None}, synchronize_session=False
        )
        db.query(Intern).filter(Intern.mentor_id == account.id).update(
            {Intern.mentor_id: None}, synchronize_session=False
        )

    record_event(
        db,
        "intern",
        intern.id,
        "intern.deleted",
        previous_status=previous_status,
        new_status="deleted",
        actor_id=user.id,
        meta={"name": name, "email": email},
    )
    db.delete(intern)
    db.flush()
    if account:
        db.delete(account)
    db.commit()
    return {"message": f"{name or 'Intern'} was deleted", "deleted_id": intern_id}


@router.get("/{intern_id}")
def get_intern(
    intern_id: int,
    user: User = Depends(require_roles("admin", "mentor")),
    db: Session = Depends(get_db),
) -> dict:
    intern = db.get(Intern, intern_id)
    if not intern:
        raise HTTPException(status_code=404, detail="Intern not found")
    if user.role == "mentor" and intern.mentor_id != user.id:
        raise HTTPException(status_code=403, detail="This intern is not assigned to you")
    return intern_dict(intern)


@router.patch("/{intern_id}")
def update_intern(
    intern_id: int,
    payload: InternUpdate,
    user: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    intern = db.get(Intern, intern_id)
    if not intern:
        raise HTTPException(status_code=404, detail="Intern not found")
    if payload.mentor_id is not None:
        mentor = _mentor_or_404(db, payload.mentor_id)
        intern.mentor_id = mentor.id if mentor else None
    if payload.name is not None:
        intern.user.name = payload.name
    if payload.batch is not None:
        intern.batch = payload.batch
    if payload.domain is not None:
        intern.domain = payload.domain
    if payload.skills is not None:
        intern.skills = payload.skills
    if payload.active is not None:
        intern.active = payload.active
        intern.user.active = payload.active
    record_event(
        db,
        "intern",
        intern.id,
        "intern.updated",
        actor_id=user.id,
        meta=payload.model_dump(exclude_unset=True),
    )
    db.commit()
    db.refresh(intern)
    return intern_dict(intern)
